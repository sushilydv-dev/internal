"""
Extracts text AND hyperlinks out of resume files (PDF / DOCX) so both can be
handed to the AI scorer. Works directly on in-memory bytes - nothing is
written to disk, since resumes are streamed straight from WorkDrive.

Why hyperlinks matter: resumes very often show a word like "LinkedIn" or
"GitHub" that is clickable, with the real URL stored only in a link
annotation. Plain text extraction returns just the word, so the URL was
invisible to the AI and never reached Recruit.
"""
import io
import re
import pdfplumber
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT


class UnsupportedResumeFormat(Exception):
    pass


_URL_IN_TEXT = re.compile(
    r"(?:https?://|www\.)[^\s<>()\"']+|"
    r"\b(?:linkedin\.com|github\.com|gitlab\.com|behance\.net|dribbble\.com|"
    r"leetcode\.com|hackerrank\.com|codechef\.com|codeforces\.com|kaggle\.com|"
    r"medium\.com|stackoverflow\.com)/[^\s<>()\"']+",
    re.IGNORECASE,
)


def extract(file_bytes: bytes, filename: str) -> dict:
    """Returns {"text": str, "links": [str]}; links are de-duplicated and
    include both clickable annotations and URLs written out in the text."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        text, links = _extract_pdf(file_bytes)
    elif ext == "docx":
        text, links = _extract_docx(file_bytes)
    else:
        raise UnsupportedResumeFormat(
            f"'{filename}' is a .{ext} file - only .pdf and .docx resumes are supported."
        )

    links.extend(m.group(0).rstrip(".,;:") for m in _URL_IN_TEXT.finditer(text))
    return {"text": text, "links": _dedupe_links(links)}


def extract_text(file_bytes: bytes, filename: str) -> str:
    return extract(file_bytes, filename)["text"]


def build_ai_input(parsed: dict) -> str:
    """Text for the AI, with every link appended in its own section so hidden
    URLs are visible to the model."""
    text = parsed["text"]
    if parsed["links"]:
        text += "\n\n--- HYPERLINKS FOUND IN THE DOCUMENT (including hidden clickable links) ---\n"
        text += "\n".join(parsed["links"])
    return text


def _extract_pdf(file_bytes: bytes) -> tuple[str, list]:
    text_parts, links = [], []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
            for link in page.hyperlinks or []:
                if link.get("uri"):
                    links.append(link["uri"])
    return "\n".join(text_parts).strip(), links


def _extract_docx(file_bytes: bytes) -> tuple[str, list]:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    for section in doc.sections:
        for part in (section.header, section.footer):
            parts.extend(p.text for p in part.paragraphs if p.text)

    links = [
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.reltype == RT.HYPERLINK and rel.is_external
    ]
    return "\n".join(parts).strip(), links


def _dedupe_links(links: list) -> list:
    seen, out = set(), []
    for link in links:
        link = (link or "").strip()
        if not link:
            continue
        key = re.sub(r"^(https?://)?(www\.)?", "", link.lower()).rstrip("/")
        if key in seen:
            continue
        seen.add(key)
        out.append(link)
    return out
